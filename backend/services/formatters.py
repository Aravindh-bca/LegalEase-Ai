from io import BytesIO
from docx import Document
from fpdf import FPDF


def format_txt(content: str) -> bytes:
    return content.encode("utf-8")


def format_docx(content: str) -> bytes:
    document = Document()

    for line in content.splitlines():
        document.add_paragraph(line)

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output.getvalue()


def format_pdf(content: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=11)

    safe_content = content.encode(
        "latin-1",
        "replace"
    ).decode("latin-1")

    for line in safe_content.splitlines():

        line = line.strip()

        if not line:
            pdf.ln(6)
            continue

        while len(line) > 70:

            part = line[:70]
            space = part.rfind(" ")

            if space > 0:
                part = part[:space]

            pdf.cell(
                0,
                7,
                part,
                new_x="LMARGIN",
                new_y="NEXT"
            )

            line = line[len(part):].strip()

        if line:
            pdf.cell(
                0,
                7,
                line,
                new_x="LMARGIN",
                new_y="NEXT"
            )

    return bytes(pdf.output())


def format_html_preview(content: str) -> str:
    html = content.replace("&", "&amp;")
    html = html.replace("<", "&lt;")
    html = html.replace(">", "&gt;")
    html = html.replace("\n", "<br>")

    return f"""
    <div style="
        padding: 20px;
        border: 1px solid #ccc;
        border-radius: 10px;
        background-color: white;
        color: black;
        font-family: Arial;
        line-height: 1.6;
    ">
        {html}
    </div>
    """